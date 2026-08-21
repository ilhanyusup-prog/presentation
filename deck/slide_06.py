"""Slide 6 - what's included, part two.

Light page: three icon groups, each a mint tile with a line art glyph beside a
heading and three body lines, separated by full width hairlines. Below them the
energy path diagram: five dashed circles joined by dashed connectors.
"""

from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn

from .kit import (
    ASSETS,
    CONTENT_R,
    GREEN,
    H1_SIZE,
    INK,
    LINE,
    MARGIN,
    MUTED,
    TINT,
    TRACK_LABEL,
    add_oval,
    add_picture,
    add_rule,
    add_text,
    chrome,
    fill_solid,
    lines,
    mid,
    outline,
    plain,
    u,
)

TRACK_H1 = -0.0093  # the source tightens the headline by half a point per gap

TILE_D = 85.5  # icon tile, both in the list and in the diagram
GLYPH = 42.0  # square the cropped line art sits in, centred on the tile

GROUP_STEP = 253.5  # the three list groups repeat on this pitch
TILE_CENTER = (102.75, 281.25)
TEXT_X = 169.5
HEADING_SPAN = (239.8, 273.4)
BODY_TOP_SPAN = (287.3, 314.1)
BODY_STEP = 36.75
BODY_LIFT = 4.2  # exact line spacing sinks a multi line block in the renderers
RULE_Y = 441.0

GROUPS = (
    (
        "Mounting structures",
        "icon_p06_mounting.png",
        (
            "Aluminium 6063 T6, anodized",
            "10-year product warranty",
            "Flat roof, tile roof or carport frames",
        ),
    ),
    (
        "Switchgear and connection",
        "icon_p06_switchgear.png",
        (
            "ABB, Schneider or Eaton components",
            "Fully compliant with utility regulations",
            "Smart bidirectional meter installed",
        ),
    ),
    (
        "Installation and workmanship",
        "icon_p06_installation.png",
        (
            "Our own in-house mechanical and electrical teams",
            "Crane lifting and scaffolding included",
            "Testing, commissioning and handover",
        ),
    ),
)

RING_D = 103.0  # dashed ring, measured on its stroke centre line
CAPTION_DROP = 75.95  # caption centre below the circle centre

# label, icon asset, circle centre x, circle centre y, caption bbox x0
STOPS = (
    ("Sun", "icon_p06_sun.png", 112.5, 1099.5, 93.8),
    ("Array", "icon_p06_array.png", 297.75, 1099.5, 272.7),
    ("Inverter", "icon_p06_inverter.png", 482.25, 1099.5, 446.6),
    ("Villa", "icon_p06_villa.png", 667.5, 1099.5, 647.8),
    ("Battery", "icon_p06_battery.png", 482.25, 1273.5, 449.0),
)

CONNECTORS = (
    ("Connector - sun to array", 165.0, 1099.5, 245.25, 1099.5),
    ("Connector - array to inverter", 350.25, 1099.5, 429.75, 1099.5),
    ("Connector - inverter to villa", 534.75, 1099.5, 615.0, 1099.5),
    ("Connector - inverter to battery", 482.25, 1152.0, 482.25, 1221.0),
)

FLOW_LABELS = (
    ("DC", 374.8, (1066.2, 1089.7)),
    ("AC", 560.3, (1066.2, 1089.7)),
    ("stored DC", 496.5, (1199.7, 1223.2)),
)


def flatten(slide):
    """Drop the theme style reference python-pptx puts on every autoshape.

    Every shape on this page is formatted explicitly, and the inherited
    `effectRef` and `lnRef` make renderers other than PowerPoint itself add a
    drop shadow and ignore the dashed strokes.
    """
    for shape in slide.shapes:
        for child in list(shape._element):
            if child.tag == qn("p:style"):
                shape._element.remove(child)


def dashed_line(slide, name, x0, y0, x1, y1, color, width, dash):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, u(x0), u(y0), u(x1), u(y1)
    )
    connector.name = name
    connector.line.color.rgb = color
    connector.line.width = u(width)
    connector.line.dash_style = dash
    return connector


def add_tile(slide, name, center_x, center_y, asset):
    fill_solid(add_oval(slide, f"{name} tile", center_x, center_y, TILE_D), TINT)
    return add_picture(
        slide,
        f"{name} icon",
        ASSETS / asset,
        center_x - GLYPH / 2,
        center_y - GLYPH / 2,
        GLYPH,
        GLYPH,
    )


def build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    chrome(slide, 6, "WHAT'S INCLUDED \u00b7 2/2")
    add_text(
        slide,
        "Headline",
        MARGIN,
        mid(156.8, 202.1),
        CONTENT_R - MARGIN + 40,
        [plain("What's included")],
        type_size=H1_SIZE,
        bold=True,
        tracking=TRACK_H1,
    )

    for index, (heading, asset, body) in enumerate(GROUPS):
        offset = index * GROUP_STEP
        add_tile(slide, heading, TILE_CENTER[0], TILE_CENTER[1] + offset, asset)
        add_text(
            slide,
            f"{heading} heading",
            TEXT_X,
            mid(*HEADING_SPAN) + offset,
            CONTENT_R - TEXT_X + 60,
            [plain(heading)],
            type_size=30,
            bold=True,
        )
        add_text(
            slide,
            f"{heading} body",
            TEXT_X,
            mid(BODY_TOP_SPAN[0], BODY_TOP_SPAN[1] + 2 * BODY_STEP)
            + offset
            - BODY_LIFT,
            CONTENT_R - TEXT_X + 60,
            lines(*body),
            type_size=24,
            color=MUTED,
            line_height=BODY_STEP,
        )
        add_rule(
            slide,
            f"Rule below {heading.lower()}",
            MARGIN,
            RULE_Y + offset,
            CONTENT_R - MARGIN,
        )

    add_text(
        slide,
        "Energy path label",
        MARGIN,
        mid(998.7, 1022.2),
        400,
        [plain("ENERGY PATH")],
        type_size=21,
        bold=True,
        color=GREEN,
        tracking=TRACK_LABEL,
    )

    for name, x0, y0, x1, y1 in CONNECTORS:
        dashed_line(slide, name, x0, y0, x1, y1, GREEN, 1.5, MSO_LINE_DASH_STYLE.DASH)

    for label, asset, center_x, center_y, caption_x in STOPS:
        ring = add_oval(slide, f"{label} ring", center_x, center_y, RING_D)
        outline(ring, LINE, 1.5)
        ring.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        add_tile(slide, label, center_x, center_y, asset)
        add_text(
            slide,
            f"{label} caption",
            caption_x,
            center_y + CAPTION_DROP,
            200,
            [plain(label)],
            type_size=21,
            color=INK,
        )

    for text, left, span in FLOW_LABELS:
        add_text(
            slide,
            f"Flow label - {text}",
            left,
            mid(*span),
            200,
            [plain(text)],
            type_size=21,
            color=MUTED,
        )

    flatten(slide)
    return slide
