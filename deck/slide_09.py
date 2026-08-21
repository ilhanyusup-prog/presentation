"""Slide 9 - the team.

Light page: eyebrow and headline, a thin rail running the height of the content
with a dark ink to mint gradient, and four groups beside it, each a round icon
tile, a heading and a placeholder paragraph.
"""

from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from .kit import (
    ASSETS,
    INK,
    MINT,
    MINT_DEEP,
    MINT_SOFT,
    MUTED,
    TINT,
    add_h1,
    add_oval,
    add_picture,
    add_pill,
    add_text,
    chrome,
    fill_solid,
    lines,
    mid,
    plain,
)

# The rail runs darker for its first half than a straight ink to mint ramp
# would, so it carries this measured stop in between.
RAIL_DEEP = RGBColor(0x0A, 0x59, 0x41)

RAIL_BOX = (60.0, 238.0, 5.0, 1127.0)
RAIL_STOPS = (
    (0.00, INK),
    (0.31, RAIL_DEEP),
    (0.62, MINT_DEEP),
    (0.89, MINT),
    (1.00, MINT_SOFT),
)

TILE_X = 140.25  # centre of the icon tiles
TILE_D = 85.5
GLYPH_D = 54.0  # the cropped line art, well inside the tile
TEXT_X = 207.0
HEADING_LEADING = 36.0
BODY = "[paragraph \u00b7 from source]"

# label, icon, tile centre, heading lines, heading span, paragraph span
GROUPS = (
    (
        "Engineers",
        "sun",
        281.25,
        ("Engineers, not installers",),
        (239.8, 273.4),
        (288.0, 314.8),
    ),
    (
        "Language",
        "pulse",
        612.0,
        ("We speak your language",),
        (570.6, 604.1),
        (618.8, 645.6),
    ),
    (
        "Schedule",
        "clock",
        942.75,
        ("Nothing rushed, nothing left to", "chance"),
        (901.3, 970.9),
        (985.5, 1012.3),
    ),
    (
        "Warranty",
        "shield",
        1304.25,
        ("Five years on our workmanship",),
        (1262.8, 1296.4),
        (1311.0, 1337.8),
    ),
)


def fill_gradient_stops(shape, stops, angle=90):
    """Linear gradient with more stops than the kit's two.

    `stops` is a sequence of (position 0.0 - 1.0, colour) and `angle` is in
    degrees clockwise from left to right, so 90 runs down the page.
    """
    shape.fill.gradient()
    gradient = shape.fill._xPr.find(qn("a:gradFill"))
    stop_list = gradient.find(qn("a:gsLst"))
    for child in list(stop_list):
        stop_list.remove(child)
    for position, color in stops:
        stop_list.append(
            parse_xml(
                '<a:gs %s pos="%d"><a:srgbClr val="%s"/></a:gs>'
                % (nsdecls("a"), round(position * 100000), color)
            )
        )
    gradient.find(qn("a:lin")).set("ang", str(round(angle * 60000)))
    return shape


def build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    chrome(slide, 9, "THE TEAM")

    add_h1(slide, "Who builds your system", (156.8, 202.1))

    fill_gradient_stops(add_pill(slide, "Team rail", *RAIL_BOX), RAIL_STOPS)

    for label, icon, center_y, heading, heading_span, body_span in GROUPS:
        fill_solid(add_oval(slide, f"{label} tile", TILE_X, center_y, TILE_D), TINT)
        add_picture(
            slide,
            f"{label} icon",
            ASSETS / f"icon_p09_{icon}.png",
            TILE_X - GLYPH_D / 2,
            center_y - GLYPH_D / 2,
            GLYPH_D,
            GLYPH_D,
        )
        add_text(
            slide,
            f"{label} heading",
            TEXT_X,
            mid(*heading_span),
            500,
            lines(*heading),
            type_size=30,
            bold=True,
            line_height=HEADING_LEADING,
        )
        add_text(
            slide,
            f"{label} paragraph",
            TEXT_X,
            mid(*body_span),
            300,
            [plain(BODY)],
            type_size=24,
            color=MUTED,
        )
    return slide
