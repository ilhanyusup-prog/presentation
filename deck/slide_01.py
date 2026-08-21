"""Slide 1 - cover.

Dark page: the villa render bleeds across it under an ink gradient, the logo
sits centred at the top, and the eyebrow, headline and two meta columns stack
against the bottom edge.
"""

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from .kit import (
    ARIAL_BLACK,
    ASSETS,
    DOT_D,
    INK,
    LINE,
    MINT,
    MINT_DEEP,
    RAIL_H,
    RAIL_TOP,
    RAIL_W,
    RAIL_X,
    REF_H,
    SLIDE_COUNT,
    TRACK_DISPLAY,
    TRACK_LABEL,
    WHITE,
    add_background,
    add_logo,
    add_oval,
    add_picture,
    add_pill,
    add_shape,
    add_text,
    fill_gradient,
    fill_solid,
    mid,
    plain,
    veiled,
)

# Greys of the cover type, sampled off the reference page.
EYEBROW_GREY = RGBColor(0xA3, 0xAB, 0xAA)
LABEL_GREY = RGBColor(0x85, 0x92, 0x90)
BODY_GREY = RGBColor(0xAA, 0xB3, 0xB1)
NUMBER_GREY = RGBColor(0xBC, 0xC2, 0xC1)
RAIL_DEEP = RGBColor(0x08, 0x38, 0x2A)

# Ink veil over the hero render, sampled off the reference page as
# (fraction of page height, opacity of the render).
HERO_VEIL = (
    (0.000, 0.54),
    (0.042, 0.75),
    (0.083, 0.77),
    (0.125, 0.82),
    (0.167, 0.89),
    (0.208, 0.94),
    (0.250, 0.97),
    (0.292, 1.00),
    (0.542, 1.00),
    (0.583, 0.97),
    (0.625, 0.83),
    (0.667, 0.76),
    (0.708, 0.64),
    (0.750, 0.43),
    (0.792, 0.29),
    (0.833, 0.21),
    (0.875, 0.20),
    (0.917, 0.07),
    (0.958, 0.00),
    (1.000, 0.00),
)

COL_X = 48.0  # left margin of the cover text column
COL2_X = 411.8  # second meta column
RAIL_FILL_H = RAIL_H / SLIDE_COUNT
DOT_X, DOT_Y = RAIL_X + RAIL_W / 2, RAIL_TOP + RAIL_FILL_H


def build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    add_background(slide, INK)

    # The render bleeds a fraction of a point past both side edges, as it does
    # on the reference page, so its own aspect ratio is preserved.
    add_picture(
        slide,
        "Hero render - villa with solar array",
        veiled(ASSETS / "hero_villa.png", HERO_VEIL),
        -0.21,
        0,
        810.75,
        REF_H,
    )

    add_logo(slide, box=(339.75, 75, 130.5, 105))

    fill_gradient(
        add_pill(slide, "Accent dash", COL_X, 1020, 42, 6), MINT_DEEP, MINT, angle=0
    )
    add_text(
        slide,
        "Eyebrow - reference",
        108,
        mid(1010.7, 1034.2),
        520,
        [plain("SOLAR PROPOSAL \u00b7 NFX-[0000-0000]")],
        type_size=21,
        bold=True,
        color=EYEBROW_GREY,
        tracking=TRACK_LABEL,
    )

    add_text(
        slide,
        "Headline",
        COL_X,
        mid(1050.0, 1232.2) + 6,  # optical centre of the two line block
        640,
        [plain("Power your villa"), plain("with solar")],
        type_size=75,
        font=ARIAL_BLACK,
        color=WHITE,
        tracking=TRACK_DISPLAY,
        line_height=76.5,
    )

    add_text(
        slide,
        "Label - prepared for",
        COL_X,
        mid(1271.7, 1295.2),
        320,
        [plain("PREPARED FOR")],
        type_size=21,
        bold=True,
        color=LABEL_GREY,
        tracking=TRACK_LABEL,
    )
    add_text(
        slide,
        "Client name",
        COL_X,
        mid(1303.3, 1336.9),
        320,
        [plain("Mr Ahmed")],
        type_size=30,
        bold=True,
        color=WHITE,
    )
    add_text(
        slide,
        "Client community",
        COL_X,
        mid(1344.8, 1371.6),
        320,
        [plain("Jumeirah Golf Estates")],
        type_size=24,
        color=BODY_GREY,
    )
    add_text(
        slide,
        "Proposal date",
        COL_X,
        mid(1390.5, 1417.3),
        320,
        [plain("Aug 19, 2026")],
        type_size=24,
        color=LABEL_GREY,
    )

    fill_solid(
        add_shape(slide, "Meta divider", 377.25, 1272, 1.5, 154.5), WHITE, opacity=0.28
    )

    add_text(
        slide,
        "Label - system size",
        COL2_X,
        mid(1271.7, 1295.2),
        290,
        [plain("SYSTEM SIZE")],
        type_size=21,
        bold=True,
        color=LABEL_GREY,
        tracking=TRACK_LABEL,
    )
    add_text(
        slide,
        "System size",
        COL2_X,
        mid(1307.1, 1423.1) - 4,  # optical centre of the three line block
        290,
        [
            plain("12 kW system \u00b7"),
            [
                ("covers ", {}),
                ("84%", {"font": ARIAL_BLACK, "color": MINT}),
                (" of your", {}),
            ],
            plain("home"),
        ],
        type_size=30,
        line_height=41,
    )

    fill_solid(
        add_pill(slide, "Progress rail track", RAIL_X, RAIL_TOP, RAIL_W, RAIL_H),
        LINE,
        opacity=0.30,
    )
    fill_gradient(
        add_pill(slide, "Progress rail fill", RAIL_X, RAIL_TOP, RAIL_W, RAIL_FILL_H),
        RAIL_DEEP,
        MINT,
        angle=270,
    )
    for scale, opacity in ((2.1, 0.10), (1.55, 0.20)):
        fill_solid(
            add_oval(
                slide, f"Progress dot glow {scale:g}x", DOT_X, DOT_Y, DOT_D * scale
            ),
            MINT,
            opacity=opacity,
        )
    fill_solid(add_oval(slide, "Progress dot", DOT_X, DOT_Y, DOT_D), MINT)

    add_text(
        slide,
        "Page number",
        657.5,
        mid(160.2, 183.7),
        100,
        [plain("01")],
        type_size=21,
        bold=True,
        color=NUMBER_GREY,
        align=PP_ALIGN.RIGHT,
    )
    return slide
