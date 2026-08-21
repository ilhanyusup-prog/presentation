"""Slide 12 - monitoring.

Light page: headline and intro paragraph, then a column of four icon rows
beside the tall app screenshot slot, and a soft mint card at the foot of the
page carrying the store download call to action.
"""

from pptx.dml.color import RGBColor

from .kit import (
    ASSETS,
    GREEN_DEEP,
    H1_SIZE,
    MARGIN,
    MINT_SOFT,
    MUTED,
    TINT,
    add_oval,
    add_picture,
    add_pill,
    add_round,
    add_text,
    chrome,
    fill_solid,
    image_slot,
    lines,
    mid,
    plain,
)

# Empty picture frames sitting on the mint card, not on the paper.
SLOT_ON_CARD = RGBColor(0xE8, 0xF0, 0xEC)

# The reference sets its 40.5 pt headings a touch tight.
TRACK_HEADING = -0.009

# A text block is centred on its layout line boxes, which sit a twentieth of the
# type size below the font box the reference is measured from, so every measured
# centre is lifted by that much.
LIFT = 0.05

TILE_D = 85.5  # icon tile diameter
ICON_D = 40.0  # line art crop centred in the tile
LABEL_X = 163.5

# Feature rows: tile top, icon crop, label lines.
FEATURES = (
    (387.0, "icon_p12_statistics.png", ("Energy consumption", "statistics")),
    (490.5, "icon_p12_smart.png", ("Smart management",)),
    (594.0, "icon_p12_forecast.png", ("Generation forecast",)),
    (697.5, "icon_p12_monitoring.png", ("Live online", "monitoring")),
)

# Store rows: QR slot left, pill left, pill width, label left, store.
STORES = (
    (93.0, 93.75, 133.5, 110.3, "App Store"),
    (261.75, 252.0, 154.5, 268.5, "Google Play"),
)


def build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    chrome(slide, 12, "MONITORING")

    add_text(
        slide,
        "Headline",
        MARGIN,
        mid(150.8, 196.1) - LIFT * H1_SIZE,
        700,
        [plain("Your system, in your pocket.")],
        type_size=H1_SIZE,
        bold=True,
        tracking=TRACK_HEADING,
    )
    add_text(
        slide,
        "Intro paragraph",
        MARGIN,
        mid(231.6, 352.1) - 5.25,  # optical centre of the three line block
        620,
        lines(
            "Smart control with built-in automation \u2014 see",
            "generation, consumption and system health",
            "from anywhere.",
        ),
        type_size=30,
        line_height=43.5,
    )

    for index, (top, icon, label) in enumerate(FEATURES, start=1):
        center_y = top + TILE_D / 2
        fill_solid(
            add_oval(
                slide,
                f"Feature row {index} tile",
                MARGIN + TILE_D / 2,
                center_y,
                TILE_D,
            ),
            TINT,
        )
        add_picture(
            slide,
            f"Feature row {index} icon",
            ASSETS / icon,
            MARGIN + (TILE_D - ICON_D) / 2,
            center_y - ICON_D / 2,
            ICON_D,
            ICON_D,
        )
        add_text(
            slide,
            f"Feature row {index} label",
            LABEL_X,
            center_y - LIFT * 30,
            400,
            lines(*label),
            type_size=30,
            bold=True,
            line_height=36,
        )

    image_slot(slide, "IMG-14", 495, 387, 225, 393.75, "IMG-14", radius=21)

    fill_solid(add_round(slide, "Download card", MARGIN, 969, 660, 396, 24), TINT)
    add_text(
        slide,
        "Card heading",
        93,
        mid(1002.8, 1048.1) - LIFT * H1_SIZE,
        600,
        [plain("Ready to catch the sun?")],
        type_size=H1_SIZE,
        bold=True,
        tracking=TRACK_HEADING,
    )
    add_text(
        slide,
        "Card paragraph",
        93,
        mid(1063.5, 1124.1) - 3.75,  # optical centre of the two line block
        600,
        lines(
            "Download the app and follow your villa's energy from",
            "day one.",
        ),
        type_size=24,
        color=MUTED,
        line_height=33.8,
    )

    for slot_left, pill_left, pill_width, label_left, store in STORES:
        image_slot(
            slide,
            f"{store} QR",
            slot_left,
            1152,
            135,
            135,
            f"{store} QR",
            radius=21,
            tint=SLOT_ON_CARD,
        )
        fill_solid(
            add_pill(slide, f"{store} button", pill_left, 1299, pill_width, 33),
            MINT_SOFT,
        )
        add_text(
            slide,
            f"{store} button label",
            label_left,
            mid(1303.2, 1326.7) - LIFT * 21,
            pill_width,
            [plain(store)],
            type_size=21,
            bold=True,
            color=GREEN_DEEP,
        )

    image_slot(
        slide,
        "Deye logo",
        507,
        1257,
        180,
        75,
        "Deye logo",
        radius=15,
        tint=SLOT_ON_CARD,
    )
    return slide
