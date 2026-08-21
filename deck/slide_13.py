"""Slide 13 - closing.

The second dark page: an ink field carrying the full bleed closing photo, the
wordmark at the top left and, stacked against the bottom edge, the display
headline, the WhatsApp button and the contact block.
"""

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from .kit import (
    ARIAL_BLACK,
    ASSETS,
    GREEN_DEEP,
    INK,
    LINE,
    MARGIN,
    MINT,
    RAIL_TOP,
    RAIL_W,
    RAIL_X,
    REF_H,
    REF_W,
    TRACK_DISPLAY,
    WHITE,
    add_background,
    add_page_number,
    add_picture,
    add_pill,
    add_round,
    add_text,
    fill_gradient,
    fill_solid,
    image_slot,
    lines,
    mid,
    plain,
)

# Sampled off the reference page: the empty photo frame over the ink field, the
# light end of the button gradient, and the greys of the closing type.
SLOT_DARK = RGBColor(0x08, 0x27, 0x1E)
BUTTON_LIGHT = RGBColor(0x15, 0xB8, 0x8C)
ADDRESS_GREY = RGBColor(0xA9, 0xB4, 0xB1)
NUMBER_GREY = RGBColor(0xB5, 0xBE, 0xBC)

# The reference tracks its uppercase address out, and sets the button label a
# touch tighter than the metric width.
TRACK_ADDRESS = 0.035
TRACK_BUTTON = -0.0095

# The closing page carries the full wordmark, wider than the inner page mark.
LOGO_PANEL = (60.0, 75.0, 93.75, 75.0)
LOGO_ART = (60.5, 75.0, 93.0, 75.0)

RAIL_H = 1058.0  # the closing page fills the rail down to the foot of the button
BUTTON = (MARGIN, 1132.0, 717.0, 106.0)


def build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    add_background(slide, INK)
    image_slot(
        slide,
        "Closing photo",
        0,
        0,
        REF_W,
        REF_H,
        "Closing photo \u2014 villa at night, lit",
        radius=0,
        tint=SLOT_DARK,
        dark=True,
    )

    fill_solid(
        add_round(slide, "Logo panel", *LOGO_PANEL, 10.5),
        RGBColor(0x7F, 0x7F, 0x7F),
        opacity=0.078,
    )
    add_picture(
        slide,
        "Almuhib Solar Energy logo",
        ASSETS / "logo_apm.png",
        *LOGO_ART,
    )

    # The rail is full on the last page, so the track sits entirely under it.
    fill_solid(
        add_pill(slide, "Progress rail track", RAIL_X, RAIL_TOP, RAIL_W, RAIL_H),
        LINE,
        opacity=0.30,
    )
    fill_gradient(
        add_pill(slide, "Progress rail fill", RAIL_X, RAIL_TOP, RAIL_W, RAIL_H),
        INK,
        MINT,
        angle=270,
    )
    add_page_number(slide, 13, mid(1004.7, 1028.2), color=NUMBER_GREY)

    add_text(
        slide,
        "Headline",
        MARGIN,
        mid(834.0, 1092.7) + 6,  # optical centre of the three line block
        640,
        lines("YOUR OWN", "POWER DAY", "AND NIGHT"),
        type_size=75,
        font=ARIAL_BLACK,
        color=WHITE,
        tracking=TRACK_DISPLAY,
        line_height=76.5,
    )

    fill_gradient(
        add_round(slide, "WhatsApp button", *BUTTON, 24), GREEN_DEEP, BUTTON_LIGHT
    )
    add_text(
        slide,
        "WhatsApp button label",
        BUTTON[0],
        mid(1161.4, 1208.0),
        BUTTON[2],
        [plain("Message us on WhatsApp")],
        type_size=33,
        font=ARIAL_BLACK,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        tracking=TRACK_BUTTON,
    )

    add_text(
        slide,
        "Phone number",
        MARGIN,
        mid(1261.3, 1294.9) - 1.5,  # optical centre of the line
        400,
        [plain("+971 58 583 7425")],
        type_size=30,
        bold=True,
        color=WHITE,
    )
    add_text(
        slide,
        "Address",
        MARGIN,
        mid(1307.0, 1359.7) - 3.2,  # optical centre of the two line block
        620,
        lines(
            "OFFICE 101-34, GAL BUSINESS CENTER, DUBAI",
            "PRODUCTION CITY, DUBAI, UAE",
        ),
        type_size=21,
        color=ADDRESS_GREY,
        tracking=TRACK_ADDRESS,
        line_height=29.2,
    )
    return slide
