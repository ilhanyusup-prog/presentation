"""Slide 5 - what's included, first half.

Light page: eyebrow and headline over four identical groups, each a round icon
tile at the left margin with a heading and three lines of copy beside it, and a
full width hairline between them.
"""

from pptx.oxml.ns import qn

from .kit import (
    ASSETS,
    CONTENT_R,
    H1_SIZE,
    INK,
    MARGIN,
    MUTED,
    add_picture,
    add_rule,
    add_text,
    chrome,
    lines,
    mid,
    plain,
)

# The reference tracks its page headlines in by 0.375 pt a character, which
# kit.add_h1 does not carry, and sets the line a shade low in its box.
H1_TRACK = -0.375 / 40.5
H1_LIFT = 2.1

TILE = 85.5  # icon tile diameter, at the left margin
TEXT_X = 169.5  # heading and copy column
HEAD_DY = 18.1  # heading centre, from the top of the tile
BODY_DY = 48.8  # first copy line, from the top of the tile
BODY_SIZE = 24.0
BODY_STEP = 32.2  # line height inside one wrapped item
BODY_GAP = 36.75  # from one item to the next
BODY_BOX = 26.8  # ascender to descender of a 24 pt Arial line
BODY_LIFT = 3.0  # exact line spacing sets the block this much low

GROUPS = (
    (
        238.5,
        "permits",
        "Permits and approvals",
        (
            ("DEWA, ADDC, SEWA or Etihad WE, whichever", "applies"),
            ("Community and master developer approvals",),
            ("Drawings, applications and final inspection",),
        ),
    ),
    (
        585.0,
        "panels",
        "Solar panels",
        (
            ("Jinko, JA Solar or JCL, 700 W bifacial N-type",),
            ("25-year performance warranty",),
            ("Rated for salt mist and abrasive dust",),
        ),
    ),
    (
        898.5,
        "inverter",
        "Inverter",
        (
            ("Deye on-grid or hybrid inverter",),
            ("10-year warranty",),
            ("Active cooling, holds output through peak heat",),
        ),
    ),
    (
        1212.0,
        "battery",
        "Battery storage",
        (
            ("LFP high-voltage cells",),
            ("10-year warranty",),
            ("Aerosol fire suppression inside the cabinet",),
        ),
    ),
)

RULES = (503.25, 817.5, 1131.0)


def drop_theme_style(slide):
    """Strip the theme shape style, whose effect reference paints a shadow.

    Autoshapes carry a <p:style> that points at the theme's effect list; the
    reference page draws its rules and tiles flat.
    """
    for shape in slide.shapes:
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)


def build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    chrome(slide, 5, "WHAT'S INCLUDED \u00b7 1/2")
    add_text(
        slide,
        "Headline",
        MARGIN,
        mid(156.8, 202.1) - H1_LIFT,
        CONTENT_R - MARGIN + 40,
        [plain("What's included")],
        type_size=H1_SIZE,
        bold=True,
        tracking=H1_TRACK,
    )

    for top, icon, heading, items in GROUPS:
        add_picture(
            slide,
            f"{heading} icon tile",
            ASSETS / f"icon_p05_{icon}.png",
            MARGIN,
            top,
            TILE,
            TILE,
        )
        add_text(
            slide,
            f"{heading} heading",
            TEXT_X,
            top + HEAD_DY,
            CONTENT_R - TEXT_X,
            [plain(heading)],
            type_size=30,
            bold=True,
            color=INK,
        )
        item_top = top + BODY_DY
        for index, item in enumerate(items, start=1):
            add_text(
                slide,
                f"{heading} line {index}",
                TEXT_X,
                item_top + ((len(item) - 1) * BODY_STEP + BODY_BOX) / 2 - BODY_LIFT,
                560,
                lines(*item),
                type_size=BODY_SIZE,
                color=MUTED,
                line_height=BODY_STEP,
            )
            item_top += BODY_GAP + (len(item) - 1) * BODY_STEP

    for index, top in enumerate(RULES, start=1):
        add_rule(slide, f"Group divider {index}", MARGIN, top, CONTENT_R - MARGIN)

    drop_theme_style(slide)
    return slide
