"""Slide 10 - positioning.

Light page: a six line display headline over a four line paragraph, with a row
of three feature cards along the bottom edge. The card row is rasterised in the
source document, so it is rebuilt here from shapes and live text; only the line
art glyph inside each mint tile is lifted as a picture.
"""

from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from .kit import (
    ARIAL_BLACK,
    ASSETS,
    INK,
    LINE,
    MARGIN,
    PAPER,
    TINT,
    TRACK_DISPLAY,
    add_oval,
    add_picture,
    add_round,
    add_text,
    chrome,
    fill_solid,
    lines,
    mid,
    outline,
    u,
)

CARD_TOP, CARD_W, CARD_H, CARD_PITCH = 1137.0, 208.0, 227.5, 226.0
CARD_RADIUS = 23.0
CARD_PAD = 28.25  # card edge to the icon tile and to the label
TILE_D = 85.5
TILE_CENTER_Y = 1208.0
GLYPH = 40.0

# The cards float on a very wide, very faint halo rather than a cast shadow:
# 0.06 alpha at the card edge, gone about 50 pt out.
SHADOW_BLUR, SHADOW_ALPHA = 60.0, 0.06

CARDS = (
    ("generation", ("High-power", "generation")),
    ("storage", ("Large-scale", "storage")),
    ("control", ("Smart load", "control")),
)


def soft_shadow(shape, blur, alpha):
    """Blurred, undisplaced drop shadow.

    python-pptx can only inherit a shadow from the theme, so the effect list is
    written straight onto the shape properties, replacing the empty one the kit
    leaves behind when it turns the inherited shadow off.
    """
    properties = shape.shadow._element
    for effects in properties.findall(qn("a:effectLst")):
        properties.remove(effects)
    properties.append(
        parse_xml(
            '<a:effectLst %s><a:outerShdw blurRad="%d" dist="0" dir="0" '
            'rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="%d"/>'
            "</a:srgbClr></a:outerShdw></a:effectLst>"
            % (nsdecls("a"), int(u(blur)), round(alpha * 100000))
        )
    )
    return shape


def build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    chrome(slide, 10, "POSITIONING")

    add_text(
        slide,
        "Headline",
        MARGIN,
        mid(140.96, 629.22) + 6.45,  # optical centre of the six line block
        700,
        lines(
            "First high-",
            "capacity hibrid",
            "energy system",
            "provider for",
            "luxury villas in",
            "the UAE.",
        ),
        type_size=75,
        font=ARIAL_BLACK,
        color=INK,
        tracking=TRACK_DISPLAY,
        line_height=76.5,
    )

    add_text(
        slide,
        "Paragraph",
        MARGIN,
        mid(656.09, 820.11) - 5.25,
        660,
        lines(
            "Industrial-grade engineering, applied to private",
            "villas. Conventional residential installers can't",
            "build systems of this scale \u2014 large energy",
            "companies won't.",
        ),
        type_size=30,
        color=INK,
        line_height=43.5,
    )

    for index, (glyph, label) in enumerate(CARDS):
        left = MARGIN + index * CARD_PITCH
        number = index + 1
        card = outline(
            fill_solid(
                add_round(
                    slide,
                    f"Feature card {number}",
                    left,
                    CARD_TOP,
                    CARD_W,
                    CARD_H,
                    CARD_RADIUS,
                ),
                PAPER,
            ),
            LINE,
        )
        soft_shadow(card, SHADOW_BLUR, SHADOW_ALPHA)

        tile_x = left + CARD_PAD + TILE_D / 2
        fill_solid(
            add_oval(slide, f"Icon tile {number}", tile_x, TILE_CENTER_Y, TILE_D), TINT
        )
        add_picture(
            slide,
            f"Icon {number} - {glyph}",
            ASSETS / f"icon_p10_{glyph}.png",
            tile_x - GLYPH / 2,
            TILE_CENTER_Y - GLYPH / 2,
            GLYPH,
            GLYPH,
        )
        add_text(
            slide,
            f"Feature label {number}",
            left + CARD_PAD,
            1302.8,
            CARD_W - 2 * CARD_PAD,
            lines(*label),
            type_size=27,
            bold=True,
            color=INK,
            line_height=32.4,
        )
    return slide
