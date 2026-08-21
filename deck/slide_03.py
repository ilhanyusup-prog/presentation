"""Slide 3 - cost structure.

Light page: a proportional bar splits the project into six segments that run
dark to light green, six rows spell those segments out with their spec and their
share of the spend, and a dark card at the foot carries the total.
"""

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

from .kit import (
    ARIAL_BLACK,
    CONTENT_R,
    GREEN,
    GREEN_DEEP,
    H1_SIZE,
    INK,
    MARGIN,
    MINT,
    MINT_SOFT,
    MUTED,
    TRACK_DISPLAY,
    WHITE,
    add_round,
    add_rule,
    add_shape,
    add_text,
    chrome,
    fill_solid,
    mid,
    plain,
)

TEAL = RGBColor(0x17, 0xC3, 0x9A)  # fourth step of the cost ramp

TRACK_H1 = -0.009  # the source tightens its headlines less than its display type

BAR_TOP, BAR_H, BAR_RADIUS = 238.5, 33.0, 6.0
SEGMENTS = (  # right edge, colour; the widths are the costs of the six rows
    (168.8, INK),
    (242.2, GREEN_DEEP),
    (356.2, GREEN),
    (558.0, TEAL),
    (683.2, MINT),
    (720.0, MINT_SOFT),
)

ROWS = (  # name, spec, amount
    ("Solar panels", "33 \u00d7 700 W bifacial", "18,711"),
    ("Hybrid inverter", "Deye 30 kW, on & off-grid", "12,548"),
    ("Battery storage", "4 \u00d7 20 kWh LFP", "19,560"),
    ("Installation", "Mounting, wiring, commissioning", "34,650"),
    ("Cables and components", "DC/AC cabling, breakers, connectors", "21,384"),
    ("Approvals and logistics", "Drawings, permits, transport, equipment", "6,353"),
)
ROW_PITCH = 111.75
BULLET = (357.0, 16.5, 4.5)  # top of the first bullet, side, corner radius
ROW_COL = 94.5  # text column of a row
NAME_SPAN = (347.8, 381.4)
SPEC_SPAN = (387.0, 413.8)
AMOUNT_SPAN = (348.0, 390.3)
RULE_TOP = 435.75

TOTAL_CARD = (MARGIN, 1172.2, 660.0, 192.8)
TOTAL_RADIUS = 24.0
TOTAL_COL = 93.0


def clear_theme_effects(slide):
    """Point every autoshape's theme effect reference at nothing.

    kit gives its shapes an empty <a:effectLst/>, which is what PowerPoint
    reads, but renderers that overlook it fall back on the theme effect style
    and draw its drop shadow, so the reference is cleared as well.
    """
    for shape in slide.shapes:
        style = shape._element.find(qn("p:style"))
        if style is not None:
            style.find(qn("a:effectRef")).set("idx", "0")


def build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    chrome(slide, 3, "COST STRUCTURE")
    add_text(
        slide,
        "Headline",
        MARGIN,
        mid(156.8, 202.1),
        700,
        [plain("What you're paying for")],
        type_size=H1_SIZE,
        bold=True,
        tracking=TRACK_H1,
    )

    # The bar is rounded on the outside and square where segments meet, so the
    # two end segments reach a corner radius under a neighbour drawn over them.
    edges = (MARGIN,) + tuple(right for right, _ in SEGMENTS)
    last = len(SEGMENTS) - 1
    for index in (0, last, *range(1, last)):
        left, right = edges[index], edges[index + 1]
        name = f"Bar segment {index + 1}"
        if index == 0:
            segment = add_round(
                slide, name, left, BAR_TOP, right - left + BAR_RADIUS, BAR_H, BAR_RADIUS
            )
        elif index == last:
            segment = add_round(
                slide,
                name,
                left - BAR_RADIUS,
                BAR_TOP,
                right - left + BAR_RADIUS,
                BAR_H,
                BAR_RADIUS,
            )
        else:
            segment = add_shape(slide, name, left, BAR_TOP, right - left, BAR_H)
        fill_solid(segment, SEGMENTS[index][1])

    add_text(
        slide,
        "Bar caption",
        MARGIN,
        mid(286.2, 309.7),
        620,
        [plain("Six lines, proportional to spend \u00b7 AED, VAT included")],
        type_size=21,
        color=MUTED,
    )

    bullet_top, bullet_side, bullet_radius = BULLET
    for index, ((name, spec, amount), (_, colour)) in enumerate(zip(ROWS, SEGMENTS)):
        row = index + 1
        offset = index * ROW_PITCH
        fill_solid(
            add_round(
                slide,
                f"Row {row} bullet",
                MARGIN,
                bullet_top + offset,
                bullet_side,
                bullet_side,
                bullet_radius,
            ),
            colour,
        )
        add_text(
            slide,
            f"Row {row} name",
            ROW_COL,
            mid(*NAME_SPAN) + offset,
            480,
            [plain(name)],
            type_size=30,
            bold=True,
        )
        add_text(
            slide,
            f"Row {row} spec",
            ROW_COL,
            mid(*SPEC_SPAN) + offset,
            520,
            [plain(spec)],
            type_size=24,
            color=MUTED,
        )
        add_text(
            slide,
            f"Row {row} amount",
            CONTENT_R - 300,
            mid(*AMOUNT_SPAN) + offset,
            300,
            [plain(amount)],
            type_size=30,
            font=ARIAL_BLACK,
            align=PP_ALIGN.RIGHT,
        )
        if row < len(ROWS):
            add_rule(
                slide,
                f"Row {row} rule",
                MARGIN,
                RULE_TOP + offset,
                CONTENT_R - MARGIN,
            )

    fill_solid(add_round(slide, "Total card", *TOTAL_CARD, TOTAL_RADIUS), INK)
    add_text(
        slide,
        "Total label",
        TOTAL_COL,
        mid(1205.3, 1232.1),
        520,
        [plain("Total project cost, including VAT")],
        type_size=24,
        color=WHITE,
    )
    add_text(
        slide,
        "Total amount",
        TOTAL_COL,
        mid(1239.1, 1332.2),
        560,
        [plain("AED 113,206")],
        type_size=66,
        font=ARIAL_BLACK,
        color=WHITE,
        tracking=TRACK_DISPLAY,
    )

    clear_theme_effects(slide)
    return slide
