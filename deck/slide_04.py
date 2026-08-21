"""Slide 4 - payment.

Light page: the headline sits over a dark total value card, then two option
cards list the payment stages as a mint pill plus the work it releases, and a
green accent dash opens the offer validity line at the foot of the page.
"""

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from .kit import (
    ARIAL_BLACK,
    CONTENT_R,
    GREEN,
    GREEN_DEEP,
    H1_SIZE,
    INK,
    LINE,
    MARGIN,
    MINT,
    MINT_SOFT,
    MUTED,
    TRACK_DISPLAY,
    TRACK_LABEL,
    WHITE,
    add_pill,
    add_round,
    add_rule,
    add_text,
    chrome,
    fill_gradient,
    fill_solid,
    gradient_line,
    lines,
    mid,
    outline,
    plain,
    u,
)

# The Option 2 card is stroked with a gradient, sampled at its left and right
# edges on the reference page.
BORDER_START = RGBColor(0x14, 0xAF, 0x7F)
BORDER_END = RGBColor(0x85, 0xE9, 0xAA)

# The reference tracks its page headlines in by 0.375 pt a character, which
# kit.add_h1 does not carry, and sets the line a shade low in its box.
H1_TRACK = -0.375 / 40.5
H1_LIFT = 2.1

CARD_W = CONTENT_R - MARGIN
CARD_RADIUS = 24.0
CARD_PAD = 33.0  # from the inner edge of the border to the content column

PILL_W, PILL_H = 86.3, 39.7
PCT_DX = 15.0  # percentage inset into its pill
PCT_DY = 19.75  # percentage centre, from the top of its pill
BODY_DX = 104.1  # stage copy, from the content column
BODY_DY = 1.5  # first line of the stage copy, from the top of its pill
BODY_SIZE = 24.0
BODY_STEP = 31.5  # line height inside one stage
BODY_BOX = 26.8  # ascender to descender of a 24 pt Arial line
BODY_LIFT = 3.0  # exact line spacing sets the block this much low
ROW_PAD = 8.6  # under the taller of pill and copy, before the hairline
ROW_GAP = 9.0  # under a hairline, before the next pill
RULE_H = 1.5

OPTION_1_ROWS = (
    (
        "10%",
        ("On signing \u2014 community approvals,", "documents and detailed drawings"),
    ),
    (
        "30%",
        ("Once permits and drawings are approved \u2014", "solar panels purchased"),
    ),
    ("30%", ("Panel installation, batteries and inverters", "purchased")),
    ("20%", ("System connection and electrical works",)),
    ("10%", ("Final inspection, commissioning, monitoring", "app, handover")),
)

OPTION_2_ROWS = (
    ("60%", ("On signing \u2014 approvals, drawings and full", "equipment purchase")),
    ("30%", ("Installation complete \u2014 panels mounted,", "system connected")),
    ("10%", ("Final inspection, commissioning and", "handover")),
)


def drop_theme_style(slide):
    """Strip the theme shape style, whose effect reference paints a shadow.

    Autoshapes carry a <p:style> that points at the theme's effect list; the
    reference page draws its cards and pills flat.
    """
    for shape in slide.shapes:
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)


def soft_shadow(shape, color, opacity, blur, distance):
    """Wide, low contrast drop shadow, the way the reference lifts its cards."""
    properties = shape._element.spPr
    for existing in properties.findall(qn("a:effectLst")):
        properties.remove(existing)
    properties.append(
        parse_xml(
            '<a:effectLst %s><a:outerShdw blurRad="%d" dist="%d" dir="5400000" '
            'rotWithShape="0"><a:srgbClr val="%02X%02X%02X">'
            '<a:alpha val="%d"/></a:srgbClr></a:outerShdw></a:effectLst>'
            % (nsdecls("a"), u(blur), u(distance), *color, round(opacity * 100000))
        )
    )
    return shape


def option_card(slide, label, top, height, border, title, title_y, first_pill, rows):
    """Draw one option card and its stage rows, and return the frame to stroke."""
    frame = add_round(
        slide,
        f"{label} card",
        MARGIN + border / 2,
        top + border / 2,
        CARD_W - border,
        height - border,
        CARD_RADIUS,
    )
    fill_solid(frame, WHITE)

    left = MARGIN + border + CARD_PAD
    right = CONTENT_R - border - CARD_PAD
    add_text(
        slide,
        f"{label} title",
        left,
        title_y,
        right - left,
        [plain(title)],
        type_size=21,
        bold=True,
        color=GREEN,
        tracking=TRACK_LABEL,
    )

    pill_top = first_pill
    for index, (percent, stage) in enumerate(rows, start=1):
        fill_solid(
            add_pill(
                slide, f"{label} row {index} pill", left, pill_top, PILL_W, PILL_H
            ),
            MINT_SOFT,
        )
        add_text(
            slide,
            f"{label} row {index} share",
            left + PCT_DX,
            pill_top + PCT_DY,
            PILL_W,
            [plain(percent)],
            type_size=BODY_SIZE,
            font=ARIAL_BLACK,
            color=GREEN_DEEP,
        )
        add_text(
            slide,
            f"{label} row {index} stage",
            left + BODY_DX,
            pill_top
            + BODY_DY
            + ((len(stage) - 1) * BODY_STEP + BODY_BOX) / 2
            - BODY_LIFT,
            520,
            lines(*stage),
            type_size=BODY_SIZE,
            line_height=BODY_STEP,
        )
        pill_top += max(PILL_H, len(stage) * BODY_STEP) + ROW_PAD
        if index < len(rows):
            add_rule(slide, f"{label} row {index} rule", left, pill_top, right - left)
            pill_top += RULE_H + ROW_GAP
    return frame


def build(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    chrome(slide, 4, "PAYMENT")
    add_text(
        slide,
        "Headline",
        MARGIN,
        mid(156.8, 202.1) - H1_LIFT,
        CONTENT_R - MARGIN + 40,
        [plain("Chose your payment option")],
        type_size=H1_SIZE,
        bold=True,
        tracking=H1_TRACK,
    )

    fill_solid(
        add_round(slide, "Total value card", MARGIN, 238.5, CARD_W, 121.5, CARD_RADIUS),
        INK,
    )
    add_text(
        slide,
        "Total value label",
        93.0,
        mid(285.8, 312.6),
        300,
        [plain("Total project value")],
        type_size=BODY_SIZE,
        color=WHITE,
    )
    add_text(
        slide,
        "Total value amount",
        400.6,
        mid(265.2, 332.9),
        320,
        [plain("AED 85,000")],
        type_size=48,
        font=ARIAL_BLACK,
        color=WHITE,
        tracking=TRACK_DISPLAY,
    )

    card_1 = option_card(
        slide,
        "Option 1",
        378.0,
        487.5,
        1.5,
        "OPTION 1 \u00b7 PAY AS WE BUILD",
        mid(412.3, 435.8),
        463.55,
        OPTION_1_ROWS,
    )
    outline(card_1, LINE, 1.5)
    soft_shadow(card_1, INK, 0.06, 60.0, 10.0)
    add_text(
        slide,
        "Option 1 stage count",
        485.5,
        mid(412.3, 435.8),
        200,
        [plain("5 stages")],
        type_size=21,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )

    card_2 = option_card(
        slide,
        "Option 2",
        883.5,
        359.25,
        2.25,
        "OPTION 2 \u00b7 FAST TRACK",
        mid(926.0, 949.5),
        981.8,
        OPTION_2_ROWS,
    )
    gradient_line(card_2, BORDER_START, BORDER_END, 2.25)
    fill_solid(
        add_pill(slide, "Option 2 saving pill", 543.75, 918.75, 141.0, 36.0), MINT_SOFT
    )
    add_text(
        slide,
        "Option 2 saving",
        558.6,
        mid(921.9, 951.5),
        160,
        [plain("Save 15%")],
        type_size=21,
        font=ARIAL_BLACK,
        color=GREEN_DEEP,
    )

    fill_gradient(
        add_pill(slide, "Footer accent dash", MARGIN, 1349.25, 42.0, 4.5),
        GREEN_DEEP,
        MINT,
        angle=0,
    )
    add_text(
        slide,
        "Offer validity",
        117.0,
        mid(1337.3, 1364.1),
        560,
        [plain("Offer validity \u2014 14 days \u00b7 until Aug 20, 2026.")],
        type_size=BODY_SIZE,
        color=MUTED,
    )

    drop_theme_style(slide)
    return slide
